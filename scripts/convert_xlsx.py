import json
import sys
import openpyxl
from pathlib import Path

TACTIC_SHEETS = [
    "Impact", "Collection", "Exfiltration", "Command and Control",
    "Lateral Movement", "Discovery", "Initial Access", "Execution",
    "Persistence", "Privilege Escalation", "Defense Evasion",
    "Credential Access",
]
DETECTION_SHEET = "Detection Engineering Reference"

XLSX_PATH = Path(__file__).parent.parent.parent / "PURPLEREIGN_v13_windows.xlsx"
OUT_PATH = Path(__file__).parent.parent / "src" / "data" / "purplereign.json"


def cell_val(cell):
    v = cell.value
    if v is None:
        return None
    return str(v).strip() or None


def parse_tactics(wb):
    tactics = []
    for sheet_name in TACTIC_SHEETS:
        ws = wb[sheet_name]
        test_cases = []
        for row in ws.iter_rows(min_row=2, values_only=False):
            tc_id = cell_val(row[0])
            if not tc_id:
                continue
            test_cases.append({
                "id": tc_id,
                "name": cell_val(row[1]),
                "description": cell_val(row[2]),
                "tools": cell_val(row[3]),
                "steps": cell_val(row[4]),
                "expectedDetection": cell_val(row[5]),
                "logSource": cell_val(row[6]),
                "passFail": cell_val(row[7]),
                "notes": cell_val(row[8]),
            })
        tactics.append({"name": sheet_name, "testCases": test_cases})
    return tactics


def parse_detection_reference(wb):
    ws = wb[DETECTION_SHEET]
    refs = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        tc_id = cell_val(row[0])
        if not tc_id:
            continue
        refs.append({
            "id": tc_id,
            "name": cell_val(row[1]),
            "tactics": cell_val(row[2]),
            "detectionLogic": cell_val(row[3]),
        })
    return refs


def main():
    print(f"Reading {XLSX_PATH}...")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    tactics = parse_tactics(wb)
    detection_ref = parse_detection_reference(wb)

    total = sum(len(t["testCases"]) for t in tactics)
    print(f"Parsed {total} test cases across {len(tactics)} tactics")
    print(f"Parsed {len(detection_ref)} detection reference entries")

    data = {"tactics": tactics, "detectionReference": detection_ref}
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Written to {OUT_PATH}")


if __name__ == "__main__":
    main()
